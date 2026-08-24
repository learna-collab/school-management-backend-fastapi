import secrets
from io import BytesIO
from typing import ClassVar

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.school import School
from app.models.user import User, UserRole
from app.models.user_credentials import UserCredential
from app.services.admin_service import (
    AdminService,
    generate_password,
)
from app.utils.helper import hash_password


class SchoolImportService:
    REQUIRED_COLUMNS: ClassVar[list[str]] = [
        "school_name",
        "phone",
        "whatsapp_number",
        "state",
        "address",
        "website",
        "description",
        "admin_first_name",
        "admin_last_name",
        "admin_email",
    ]

    def __init__(self):
        self.admin_service = AdminService()

    async def _generate_unique_code(
        self,
        db: AsyncSession,
    ) -> str:
        while True:
            code = secrets.token_hex(3).upper()

            result = await db.execute(select(School.id).where(School.code == code))

            existing = result.scalar_one_or_none()

            if existing is None:
                return code

    # =====================================================

    # IMPORT
    # =====================================================

    async def import_rows(
        self,
        db: AsyncSession,
        rows: list[dict],
    ):
        created = []

        try:
            for row in rows:
                # -------------------------------------
                # UNIQUE SLUG
                # -------------------------------------

                slug = await self.admin_service._generate_unique_slug(
                    db,
                    row["school_name"],
                )

                # -------------------------------------
                # UNIQUE CODE
                # -------------------------------------

                code = await self._generate_unique_code(db)

                # -------------------------------------
                # USERNAME
                # -------------------------------------

                username = await self.admin_service._generate_unique_username(
                    db,
                    row["school_name"],
                )

                # -------------------------------------
                # PASSWORD
                # -------------------------------------

                generated_password = generate_password()

                # -------------------------------------
                # CREATE SCHOOL
                # -------------------------------------

                school = School(
                    name=row["school_name"],
                    slug=slug,
                    code=code,
                    phone=row["phone"],
                    email=row["admin_email"],
                    website=row.get("website"),
                    whatsapp_number=row.get("whatsapp_number"),
                    state=row["state"],
                    address=row["address"],
                    description=row.get("description"),
                    is_active=True,
                )

                db.add(school)

                await db.flush()

                # -------------------------------------
                # CREATE ADMIN
                # -------------------------------------

                admin = User(
                    first_name=row["admin_first_name"],
                    last_name=row["admin_last_name"],
                    username=username,
                    email=row["admin_email"],
                    password_hash=hash_password(generated_password),
                    role=UserRole.SCHOOL_ADMIN,
                    school_id=school.id,
                    profile_completed=True,
                )

                db.add(admin)

                await db.flush()

                # -------------------------------------
                # SAVE CREDENTIAL
                # -------------------------------------

                credential = UserCredential(
                    school_id=school.id,
                    user_id=admin.id,
                    username=username,
                    password=generated_password,
                )

                db.add(credential)

                created.append(
                    {
                        "school_id": school.id,
                        "school_name": school.name,
                        "username": username,
                        "password": generated_password,
                    }
                )

            await db.commit()

            return created

        except Exception:
            await db.rollback()
            raise

    # =====================================================
    # READ EXCEL
    # =====================================================

    def read_excel(self, file_bytes: bytes) -> list[dict]:
        workbook = load_workbook(
            filename=BytesIO(file_bytes),
            data_only=True,
        )

        worksheet = workbook.active

        rows = list(worksheet.iter_rows(values_only=True))

        if not rows:
            raise ValueError("Excel file is empty.")

        headers = [
            str(header).strip().lower() if header is not None else ""
            for header in rows[0]
        ]

        missing_columns = [
            column for column in self.REQUIRED_COLUMNS if column not in headers
        ]

        if missing_columns:
            raise ValueError("Missing required columns: " + ", ".join(missing_columns))

        result = []

        for row in rows[1:]:
            if not any(value is not None for value in row):
                continue

            data = {}

            for index, header in enumerate(headers):
                if index >= len(row):
                    continue

                value = row[index]

                if isinstance(value, str):
                    value = value.strip()

                data[header] = value

            result.append(data)

        return result

    # =====================================================
    # VALIDATE
    # =====================================================

    async def validate_rows(
        self,
        db: AsyncSession,
        rows: list[dict],
    ):
        errors = []
        valid_rows = []

        # Track duplicates inside the uploaded file
        uploaded_emails = set()
        uploaded_slugs = set()
        uploaded_usernames = set()

        for index, row in enumerate(
            rows,
            start=2,
        ):
            row_errors = []

            school_name = row.get("school_name")

            phone = row.get("phone")

            state = row.get("state")

            address = row.get("address")

            first_name = row.get("admin_first_name")

            last_name = row.get("admin_last_name")

            email = row.get("admin_email")

            # -----------------------------------------
            # REQUIRED FIELDS
            # -----------------------------------------

            if not school_name:
                row_errors.append("School name is required.")

            if not phone:
                row_errors.append("Phone is required.")

            if not state:
                row_errors.append("State is required.")

            if not address:
                row_errors.append("Address is required.")

            if not first_name:
                row_errors.append("Admin first name is required.")

            if not last_name:
                row_errors.append("Admin last name is required.")

            if not email:
                row_errors.append("Admin email is required.")

            if row_errors:
                errors.append(
                    {
                        "row": index,
                        "errors": row_errors,
                    }
                )
                continue

            # -----------------------------------------
            # GENERATE SLUG
            # -----------------------------------------

            slug = await self.admin_service._generate_unique_slug(
                db,
                str(school_name),
            )

            # -----------------------------------------
            # CHECK EMAIL DUPLICATE
            # -----------------------------------------

            normalized_email = str(email).strip().lower()

            if normalized_email in uploaded_emails:
                row_errors.append("Duplicate admin email in uploaded file.")

            uploaded_emails.add(normalized_email)

            # -----------------------------------------
            # GENERATE USERNAME
            # -----------------------------------------

            try:
                username = await self.admin_service._generate_unique_username(
                    db,
                    str(school_name),
                )

                if username in uploaded_usernames:
                    row_errors.append(f"Duplicate generated username: {username}")

                uploaded_usernames.add(username)

            except ValueError as exc:
                row_errors.append(str(exc))
                username = None

            # -----------------------------------------
            # CHECK SLUG INSIDE FILE
            # -----------------------------------------

            if slug in uploaded_slugs:
                row_errors.append(f"Duplicate school slug: {slug}")

            uploaded_slugs.add(slug)

            # -----------------------------------------
            # FINAL RESULT
            # -----------------------------------------

            if row_errors:
                errors.append(
                    {
                        "row": index,
                        "errors": row_errors,
                    }
                )
                continue

            valid_rows.append(
                {
                    **row,
                    "school_name": str(school_name).strip(),
                    "phone": str(phone).strip(),
                    "state": str(state).strip(),
                    "address": str(address).strip(),
                    "admin_first_name": str(first_name).strip(),
                    "admin_last_name": str(last_name).strip(),
                    "admin_email": normalized_email,
                    "slug": slug,
                    "username": username,
                }
            )

        return valid_rows, errors
